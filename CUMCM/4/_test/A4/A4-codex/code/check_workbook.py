"""Independent read-only inspection of the exported Excel file."""
from pathlib import Path
import json
import numpy as np
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
def main():
 d=np.load(ROOT/'data/full_precision.npz');w=openpyxl.load_workbook(ROOT/'result4.xlsx',data_only=False)
 assert w.sheetnames==['Sheet1'];s=w.active;assert s.max_column==23 and s.max_row==len(d['time_s'])
 assert [s.cell(1,j).value for j in range(2,23)]==d['distance_cm'].tolist()
 assert s.cell(1,23).value=='药材表面'
 assert s.freeze_panes=='B2'
 count=blank=0
 for i,(t,row) in enumerate(zip(d['time_s'][1:],d['C'][1:]),2):
  assert s.cell(i,1).value==int(t)
  for j,value in enumerate(row,2):
   cell=s.cell(i,j)
   if np.isnan(value):assert cell.value is None;blank+=1
   else:assert cell.value==float(f'{value:.4f}');count+=1
   assert cell.number_format=='0.0000' and cell.data_type!='f'
 assert d['maxC'][-2]>=.15 and d['maxC'][-1]<.15
 table=np.loadtxt(ROOT/'data/table6.csv',delimiter=',',skiprows=1)
 for row in table:
  i=int(round(row[0]*3600/60));np.testing.assert_allclose(row[1:5],d['C'][i,[0,5,10,21]],rtol=0,atol=1e-15)
  assert abs(row[-1]-d['radius_m'][i]*100)<1e-14
 result=dict(records=s.max_row-1,numeric_moisture_cells=count,domain_outside_blank_cells=blank,
  fixed_distance_columns=21,surface_columns=1,format='0.0000',freeze_panes=s.freeze_panes,
  first_time_s=s.cell(2,1).value,last_time_s=s.cell(s.max_row,1).value,
  previous_maxC=float(d['maxC'][-2]),final_maxC=float(d['maxC'][-1]),
  table6_matches=True,all_cells_match=True,engine='openpyxl read-only use; no native Excel launch')
 (ROOT/'verification/workbook_check.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2))
if __name__=='__main__':main()
