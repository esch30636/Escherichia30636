"""Reopen exported workbook, independently verify every number and schema."""
from pathlib import Path
import json,hashlib
import numpy as np
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
z=np.load(ROOT/'data/full_precision.npz')
w=load_workbook(ROOT/'result2.xlsx',read_only=False,data_only=True)
assert w.sheetnames==['温度','水分浓度']
checks={}
for name,key in [('温度','T'),('水分浓度','C')]:
    s=w[name];rows=list(s.values)
    assert s.max_row==10801 and s.max_column==22
    assert rows[0][0]=='时间\\到药材中心的距离'
    assert np.allclose(rows[0][1:],np.arange(21)/10,rtol=0,atol=1e-14)
    a=np.array(rows[1:],dtype=float)
    assert np.array_equal(a[:,0],np.arange(1,10801))
    err=float(np.max(abs(a[:,1:]-np.round(z[key][1:],4))))
    assert err==0 and np.isfinite(a).all()
    assert all(s.cell(r,c).number_format=='0.0000' for r in (2,5401,10801) for c in range(2,23))
    checks[name]={'rows':10800,'radii':21,'values':226800,'max_rounding_difference':err,'range':[float(a[:,1:].min()),float(a[:,1:].max())],'freeze_panes':s.freeze_panes}
w.close()
checks['sha256']=hashlib.sha256((ROOT/'result2.xlsx').read_bytes()).hexdigest()
(ROOT/'verification/workbook_check.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
print(json.dumps(checks,ensure_ascii=False,indent=2))
