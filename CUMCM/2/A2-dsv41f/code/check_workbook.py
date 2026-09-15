"""Independent structural + numerical check of A2-gpt/result2.xlsx against
the official 附件3/result2.xlsx template and the pack's own full-precision data.

Writes verification/workbook_audit.json.
"""
from pathlib import Path
import json
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]          # A2/A2-dsv41f
A2 = ROOT.parent / "A2-gpt"                          # A2/A2-gpt
TPL = ROOT.parents[1] / "CUMCM2026Problems/A题/附件/附件3/result2.xlsx"
OUT = ROOT / "verification/workbook_audit.json"


def sheet_rows(path, name):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[name]
    rows = [r for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


rep = {}

# ---------- 1. template vs delivered: headers, indices, shape ----------
tpl_names = load_workbook(TPL, read_only=True).sheetnames
del_names = load_workbook(A2 / "result2.xlsx", read_only=True).sheetnames
rep["sheetnames"] = {"template": tpl_names, "delivered": del_names,
                     "match": tpl_names == del_names}

trows = sheet_rows(TPL, "温度")
drows = sheet_rows(A2 / "result2.xlsx", "温度")
rep["shape"] = {"template": [len(trows), len(trows[0])],
                "delivered": [len(drows), len(drows[0])]}
# Template is an abbreviated skeleton (header + 3 rows + "..." row); compare the
# parts that actually carry the required convention: header cells and radius axis.
def numeric(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)

rep["header_cell0_match"] = trows[0][0] == drows[0][0]
rep["header_cell0"] = str(trows[0][0])
rep["radius_axis_template"] = [float(v) for v in trows[0][1:] if numeric(v)]
rep["radius_axis_delivered"] = [float(v) for v in drows[0][1:] if numeric(v)]
# the template is a truncated skeleton: 0, 0.1, 0.2, "…", 2. It therefore checks
# the first three values plus the final column, not a contiguous 21-column axis.
rep["radius_axis_skeleton_match"] = bool(
    rep["radius_axis_delivered"][:3] == rep["radius_axis_template"][:3]
    and rep["radius_axis_delivered"][-1] == rep["radius_axis_template"][-1])
rep["radius_axis_step_uniform"] = bool(
    np.allclose(np.diff(rep["radius_axis_delivered"]), 0.1))
rep["radius_axis_endpoints_ok"] = (
    rep["radius_axis_delivered"][0] == 0.0
    and rep["radius_axis_delivered"][-1] == 2.0)
rep["template_time_column"] = [r[0] for r in trows[1:]]

d_time = np.array([r[0] for r in drows[1:]], dtype=float)
rep["time_index"] = {
    "start": float(d_time[0]), "end": float(d_time[-1]),
    "unique_diff": sorted(set(np.round(np.diff(d_time), 9).tolist())),
    "n_rows": int(d_time.size),
    "matches_required_grid": bool(np.array_equal(d_time, np.arange(1, 10801))),
}

# ---------- 2. delivered workbook vs full-precision arrays ----------
z = np.load(A2 / "data/full_precision.npz")
time_s = z["time_s"]; radius_cm = z["radius_cm"]
rep["full_precision_npz"] = {
    "time_s": [float(time_s[0]), float(time_s[-1]), int(time_s.size)],
    "radius_cm": [float(radius_cm[0]), float(radius_cm[-1]), int(radius_cm.size)],
}

per_sheet = {}
for name, key in (("温度", "T"), ("水分浓度", "C")):
    rows = sheet_rows(A2 / "result2.xlsx", name)
    xl = np.array([r[1:] for r in rows[1:]], dtype=float)
    full = z[key]
    assert xl.shape == (10800, 21), xl.shape
    assert full.shape == (10801, 21), full.shape
    d = xl - np.round(full[1:], 4)
    per_sheet[name] = {
        "cells": int(xl.size),
        "max_abs_diff_vs_fullprecision_rounded": float(np.max(np.abs(d))),
        "n_cells_not_equal": int(np.count_nonzero(d)),
        "delivered_min": float(xl.min()), "delivered_max": float(xl.max()),
    }
rep["delivered_vs_fullprecision"] = per_sheet

# ---------- 3. report tables 3/4 reproduce the delivered workbook ----------
def at(rows, t_sec, r_cm):
    # workbook row i (1-based) carries t = i, so index t_sec-1 for t>=1
    ti = int(np.argmin(np.abs(time_s - t_sec))) - 1
    ri = int(np.argmin(np.abs(radius_cm - r_cm)))
    return rows[ti, ri]

tab3 = [[0.5, 0, 32.1892], [0.5, 2, 35.4130], [1.0, 0, 40.3816], [1.0, 2, 42.9976],
        [2.0, 0, 48.4502], [2.0, 2, 49.0033], [3.0, 0, 49.8495], [3.0, 2, 49.9664]]
tab4 = [[0.5, 0, 2.5499], [0.5, 2, 1.6486], [1.0, 0, 2.5257], [1.0, 2, 1.4711],
        [2.0, 0, 2.1709], [2.0, 2, 1.2311], [3.0, 0, 1.7662], [3.0, 2, 1.0081]]
chk = []
for tab, key, nm in ((tab3, "T", "表3"), (tab4, "C", "表4")):
    rows = sheet_rows(A2 / "result2.xlsx", "温度" if key == "T" else "水分浓度")
    xl = np.array([r[1:] for r in rows[1:]], dtype=float)
    for h, r, v in tab:
        got = at(xl, round(h * 3600), r)
        chk.append({"table": nm, "h": h, "r_cm": r, "report": v, "workbook": got,
                    "diff": round(got - v, 6)})
rep["report_tables_vs_workbook"] = chk
rep["report_tables_max_abs_diff"] = float(max(abs(c["diff"]) for c in chk))

# ---------- 4. low-level plausibility of the delivered field ----------
T = z["T"]; C = z["C"]
rep["field_plausibility"] = {
    "T_finite": bool(np.isfinite(T).all()), "C_finite": bool(np.isfinite(C).all()),
    "T_min": float(T.min()), "T_max": float(T.max()),
    "C_min": float(C.min()), "C_max": float(C.max()),
    "T_monotone_in_r_at_end": bool(np.all(np.diff(T[-1]) >= 0)),
    "C_monotone_in_r_at_end": bool(np.all(np.diff(C[-1]) <= 0)),
    "C_negative_cells": int(np.count_nonzero(C < 0)),
    "T_below_initial_after_1h": float(T[3600:].min()),
    "T_max_minus_ambient_at_end": float(T[-1].max() - 49.5),
}
OUT.write_text(json.dumps(rep, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({k: rep[k] for k in ("sheetnames", "shape", "time_index",
                                      "delivered_vs_fullprecision",
                                      "report_tables_max_abs_diff",
                                      "field_plausibility")},
                 ensure_ascii=False, indent=2))
