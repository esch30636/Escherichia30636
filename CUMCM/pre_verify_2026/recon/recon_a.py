# -*- coding: utf-8 -*-
"""A 题数据/模板精确侦察"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import numpy as np
import openpyxl

BASE = r'D:\CUMCM\A\data'

def dump(path, maxr=5, maxc=25):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f'  [sheet] "{ws.title}"  {ws.max_row} rows x {ws.max_column} cols')
        rows = list(ws.iter_rows(values_only=True))
        print('  --- head ---')
        for i, r in enumerate(rows[:maxr]):
            cells = ['' if v is None else (f'{v:.6g}' if isinstance(v, float) else str(v))
                     for v in r[:maxc]]
            print(f'   r{i:>2}: {cells}')
        print('  --- tail ---')
        for i, r in enumerate(rows[-4:]):
            cells = ['' if v is None else (f'{v:.6g}' if isinstance(v, float) else str(v))
                     for v in r[:maxc]]
            print(f'   r{len(rows)-4+i:>2}: {cells}')
    wb.close()

print('#' * 80); print('附件1 烘房 温度/水分'); print('#' * 80)
dump(os.path.join(BASE, 'attach1_chamber_temp_moisture.xlsx'), maxr=4, maxc=5)

print()
print('#' * 80); print('附件2 半径-时间'); print('#' * 80)
dump(os.path.join(BASE, 'attach2_radius_vs_time.xlsx'), maxr=4, maxc=5)

# 附件2 数值特征
wb = openpyxl.load_workbook(os.path.join(BASE, 'attach2_radius_vs_time.xlsx'),
                            read_only=True, data_only=True)
ws = wb.worksheets[0]
rows = [r for r in ws.iter_rows(values_only=True) if r[0] is not None]
t = np.array([float(r[0]) for r in rows])
R = np.array([float(r[1]) for r in rows])
wb.close()
print(f'\n  附件2: n={len(t)}  t[0]={t[0]:.0f} t[-1]={t[-1]:.0f}  dt={t[1]-t[0]:.1f}')
print(f'        R[0]={R[0]:.6f} R[-1]={R[-1]:.6f}')
print(f'        min R={R.min():.6f} @ t={t[R.argmin()]:.0f}  (末段是否平台:')
print(f'        R[-10:]={np.array2string(R[-10:], precision=6)})')
dR = np.diff(R)
print(f'        dR/dt 首={dR[0]:.3e}  末={dR[-1]:.3e}  最小={dR.min():.3e}')
# 单次收缩率
print(f'        收缩比 R0/R_end = {R[0]/R[-1]:.6f}')

print()
print('#' * 80); print('结果模板'); print('#' * 80)
for i in (1, 2, 3, 4):
    p = os.path.join(BASE, 'templates', f'result{i}.xlsx')
    print(f'--- result{i}.xlsx ---')
    dump(p, maxr=4, maxc=25)
    print()
