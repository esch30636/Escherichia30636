# -*- coding: utf-8 -*-
"""C 题数据侦察: 列名/形状/取值/时段标签, 以及模板结构。"""
import sys, io, glob, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import numpy as np
import openpyxl

BASE = r'D:\CUMCM\C\data'

def sheet_head(path, maxr=6, maxc=10):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f'  [sheet] {ws.title}  dims={ws.max_row} x {ws.max_column}')
        for i, row in enumerate(ws.iter_rows(max_row=maxr, max_col=maxc, values_only=True)):
            cells = ['' if v is None else (f'{v:.6g}' if isinstance(v, float) else str(v))
                     for v in row]
            print(f'    r{i}: {cells}')
        # 末尾几行
        rows = list(ws.iter_rows(values_only=True))
        print(f'    ... tail:')
        for r in rows[-3:]:
            cells = ['' if v is None else (f'{v:.6g}' if isinstance(v, float) else str(v))
                     for v in r[:maxc]]
            print(f'    {cells}')
    wb.close()

for f in sorted(glob.glob(os.path.join(BASE, '*.xlsx'))):
    print('=' * 78)
    print(os.path.basename(f), f'({os.path.getsize(f)/1024:.0f} KB)')
    print('=' * 78)
    sheet_head(f)
    print()

print('#' * 78)
print('模板')
print('#' * 78)
for f in sorted(glob.glob(os.path.join(BASE, 'templates', '*.xlsx'))):
    print('=' * 78)
    print(os.path.basename(f))
    print('=' * 78)
    sheet_head(f, maxr=5, maxc=12)
    print()
