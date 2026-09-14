# -*- coding: utf-8 -*-
"""A 题: 附件2 完整半径曲线 + 结果模板结构"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import numpy as np
import openpyxl

BASE = r'D:\CUMCM\A\data'

def numeric_rows(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    out = []
    for r in ws.iter_rows(values_only=True):
        try:
            out.append([float(v) for v in r if v is not None])
        except (TypeError, ValueError):
            continue
    wb.close()
    return out

# ---------------- 附件2 ----------------
rows = numeric_rows(os.path.join(BASE, 'attach2_radius_vs_time.xlsx'))
t = np.array([r[0] for r in rows]); R = np.array([r[1] for r in rows])
print('=' * 78)
print('附件2 半径曲线')
print('=' * 78)
print(f'n={len(t)}  t: 0..{t[-1]:.0f} s @ {t[1]-t[0]:.0f} s   R: {R[0]} .. {R[-1]} cm')
print(f'\n{"t(s)":>8} {"t(h)":>8} {"R(cm)":>9} {"R/R0":>8} {"dR/dt(cm/s)":>13}')
for i in range(len(t)):
    d = f'{(R[i+1]-R[i])/(t[i+1]-t[i]):.4e}' if i + 1 < len(t) else '   -   '
    print(f'{t[i]:>8.0f} {t[i]/3600:>8.2f} {R[i]:>9.4f} {R[i]/R[0]:>8.4f} {d:>13}')

# ---------------- 模板 ----------------
print()
print('=' * 78)
print('结果模板结构')
print('=' * 78)
for i in (1, 2, 3, 4):
    p = os.path.join(BASE, 'templates', f'result{i}.xlsx')
    wb = openpyxl.load_workbook(p, data_only=True)
    print(f'\n--- result{i}.xlsx ---')
    for ws in wb.worksheets:
        print(f'  [sheet] "{ws.title}"  {ws.max_row} x {ws.max_column}   merged={len(ws.merged_cells.ranges)}')
        for ri, row in enumerate(ws.iter_rows(max_row=4, values_only=True), 1):
            cells = ['' if v is None else (f'{v:.6g}' if isinstance(v, float) else str(v))
                     for v in row[:30]]
            print(f'    r{ri}: {cells}')
        rows_all = list(ws.iter_rows(values_only=True))
        print(f'    (总行数 {len(rows_all)})')
    wb.close()
