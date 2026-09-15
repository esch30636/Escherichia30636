"""Cross-model comparison of A1-codex problem 1 against the independent reference
model (`A题建模/结果/result1.xlsx`, model v3)."""
import json
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
A1 = np.load(ROOT.parent / "A1-codex/data/full_precision.npz")
T1, C1 = A1["T_C"], A1["C_kgkg"]

wb = load_workbook(ROOT.parent / "A题建模/结果/result1.xlsx", read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
REF = {}
for name in wb.sheetnames:
    rows = list(wb[name].values)
    print(name, "dims", len(rows), len(rows[0]) if rows else 0)
    print("  header:", rows[0][:5], "...", rows[0][-2:])
    print("  first data row:", rows[1][:4])
    REF[name] = rows
wb.close()

out = {}


def grid_from(rows):
    """Return (times, radii, values) from a template sheet."""
    radii = [float(x) for x in rows[0][1:22]]
    times = []
    vals = []
    for r in rows[1:]:
        if r[0] is None or not isinstance(r[0], (int, float)):
            continue
        times.append(float(r[0]))
        vals.append([float(x) if x is not None else np.nan for x in r[1:22]])
    return np.array(times), np.array(radii), np.array(vals)


res = {}
for name, rows in REF.items():
    t, rr, v = grid_from(rows)
    print(f"\n{name}: {len(t)} times, {len(rr)} radii, t {t.min()}..{t.max()}, r {rr.min()}..{rr.max()}")
    if abs(t.min()) < 1e-9 and abs(t.max() - 1800) < 1e-9 and len(t) == 1801:
        mine = (C1 if "水分" in name or "\u6c34\u5206" in name else T1)
        # reference file includes the t=0 initial row
        ref = v
        sel = np.round(mine, 4)
        d = np.abs(sel - ref)
        res[name] = {
            "max_abs_diff": float(np.nanmax(d)),
            "n_cells_diff_gt_5e-5": int(np.sum(d > 5e-5)),
            "n_cells_diff_gt_1e-4": int(np.sum(d > 1e-4)),
            "n_cells_diff_gt_1e-3": int(np.sum(d > 1e-3)),
            "fraction_within_5e-5": float(np.mean(d < 5e-5)),
            "max": float(np.nanmax(d)),
            "max_diff_at": [int(np.unravel_index(np.nanargmax(d), d.shape)[0]),
                            int(np.unravel_index(np.nanargmax(d), d.shape)[1])],
        }
        print("  compare (my numpy vs reference xlsx, both 4dp):", res[name])
    else:
        print("  (template shape, not the full grid)")

(ROOT / "verification" / "cross_model.json").write_text(
    json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
print("\nwritten cross_model.json")
