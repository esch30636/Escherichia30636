"""Produce the independent Q1 result table (CSV) and comparison figures."""
import json
import sys
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
A1DIR = ROOT.parent / "A1-codex"
FIG = ROOT / "figures"
DATA = ROOT / "data"
VER = ROOT / "verification"
FIG.mkdir(exist_ok=True)

A1 = np.load(A1DIR / "data/full_precision.npz")
T1, C1 = A1["T_C"], A1["C_kgkg"]
ENV = np.loadtxt(A1DIR / "data/ambient.csv", delimiter=",", skiprows=1)

wb = load_workbook(ROOT.parent / "A题建模/结果/result1.xlsx", read_only=True, data_only=True)
REF = {}
for idx, name in enumerate(wb.sheetnames):
    rows = list(wb[name].values)
    REF[idx] = np.array([[float(x) for x in r[1:22]] for r in rows[1:] if isinstance(r[0], (int, float))])
wb.close()

# ---------- CSV of the independent reference model values ----------
np.savetxt(DATA / "ref_model_result1_temperature.csv", REF[0], delimiter=",",
           fmt="%.6f", header=",".join(f"r={j/10:.1f}cm" for j in range(21)), comments="")
np.savetxt(DATA / "ref_model_result1_moisture.csv", REF[1], delimiter=",",
           fmt="%.6f", header=",".join(f"r={j/10:.1f}cm" for j in range(21)), comments="")

# ---------- figures ----------
font = Path("C:/Windows/Fonts/msyh.ttc")
if font.exists():
    from matplotlib import font_manager
    font_manager.fontManager.addfont(str(font))
    plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font)).get_name()
plt.rcParams.update({"font.size": 10, "axes.unicode_minus": False, "figure.dpi": 150,
                     "savefig.dpi": 180, "axes.spines.top": False, "axes.spines.right": False})

t = np.arange(1801)
r = np.arange(21) / 10

# Fig 1: A1 vs reference model, temperature and moisture
fig, axs = plt.subplots(2, 2, figsize=(11, 7), layout="constrained")
for ax, (mine, ref, lab, unit) in zip(axs[0], [
        (T1, REF[0], "温度", "℃"), (C1, REF[1], "水分浓度", "kg/kg")]):
    for j, col in zip([0, 10, 20], ["#1d4ed8", "#059669", "#be123c"]):
        ax.plot(t, mine[:, j], color=col, lw=1.6, label=f"A1 {j/10:g} cm")
        ax.plot(t, ref[:, j], color=col, lw=1.0, ls="--", label=f"参考模型 {j/10:g} cm")
    ax.set(xlabel="时间 / s", ylabel=f"{lab} / {unit}", title=f"{lab}：A1 与独立参考模型")
    ax.grid(alpha=.15); ax.legend(fontsize=7, ncol=2)
for ax, (mine, ref, lab) in zip(axs[1], [(T1, REF[0], "温度 / ℃"), (C1, REF[1], "水分浓度 / (kg/kg)")]):
    d = np.abs(mine - ref)
    im = ax.pcolormesh(r, t, d, shading="auto", cmap="magma")
    fig.colorbar(im, ax=ax, label="|差|")
    ax.set(xlabel="径向位置 / cm", ylabel="时间 / s", title=f"{lab} 绝对差（A1 $-$ 参考模型）")
fig.savefig(FIG / "01_a1_vs_reference.png"); plt.close(fig)

# Fig 2: verification evidence
fig, axs = plt.subplots(1, 3, figsize=(13, 3.8), layout="constrained")
ax = axs[0]
for j, col in zip([0, 10, 20], ["#1d4ed8", "#059669", "#be123c"]):
    d = np.abs(T1[:, j] - REF[0][:, j])
    ax.semilogy(t, d + 1e-12, color=col, lw=1.3, label=f"r={j/10:g} cm")
ax.set(xlabel="时间 / s", ylabel="|A1 $-$ 参考| / ℃", title="温度差随时间")
ax.grid(alpha=.15, which="both"); ax.legend(fontsize=8)
ax = axs[1]
ax.semilogy(t, np.abs(C1[:, 20] - REF[1][:, 20]) + 1e-12, color="#be123c", lw=1.4)
ax.set(xlabel="时间 / s", ylabel="|A1 $-$ 参考| / (kg/kg)", title="表面水分浓度差")
ax.grid(alpha=.15, which="both")
ax = axs[2]
tN = np.linspace(0, 1800, 200)
for t_, c in [(100, "#93c5fd"), (600, "#3b82f6"), (1800, "#1e3a8a")]:
    ax.loglog([t_], [np.sqrt(np.pi * 4.9377e-9 * t_) * 1000], "o", color=c)
ax.set(xlabel="时间 / s", ylabel="水分边界层厚度 / mm", title="边界层 $\\sqrt{\\pi Dt}$ vs 表面网格")
ax.axhline(0.0125, color="#be123c", ls="--", lw=1.3);
ax.text(1, 0.014, "A1 表面单元中心距表面 0.0125 mm", color="#be123c", fontsize=8)
ax.grid(alpha=.15, which="both")
fig.savefig(FIG / "02_verification_evidence.png"); plt.close(fig)

summary = {
    "T_max_abs_diff_vs_reference_C": float(np.max(np.abs(T1 - REF[0]))),
    "C_max_abs_diff_vs_reference_kgkg": float(np.max(np.abs(C1 - REF[1]))),
    "T_final_a1": T1[-1, [0, 5, 10, 15, 20]].tolist(),
    "T_final_ref": REF[0][-1, [0, 5, 10, 15, 20]].tolist(),
    "C_final_a1": C1[-1, [0, 5, 10, 15, 20]].tolist(),
    "C_final_ref": REF[1][-1, [0, 5, 10, 15, 20]].tolist(),
}
(VER / "comparison_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
