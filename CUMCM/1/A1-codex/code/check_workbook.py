"""Read-only Excel validation, including every result and numeric cell format."""
from pathlib import Path
import json
import numpy as np
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
expected=json.loads((root/'data/workbook_values.json').read_text(encoding='utf-8'))
wb=load_workbook(root/'result1.xlsx',read_only=False,data_only=True)
assert wb.sheetnames==['温度','水分浓度']
checks={}
for name in wb.sheetnames:
    s=wb[name]
    assert (s.max_row,s.max_column)==(1801,22)
    assert s.cell(1,1).value=='时间\\到药材中心的距离'
    assert [s.cell(1,j).value for j in range(2,23)]==expected['radius_cm']
    assert [s.cell(i,1).value for i in range(2,1802)]==expected['time_s']
    vals=np.array([[s.cell(i,j).value for j in range(2,23)] for i in range(2,1802)],dtype=float)
    assert np.all(np.isfinite(vals))
    delta=float(np.max(np.abs(vals-np.array(expected[name]))))
    assert delta<5e-12
    assert all(s.cell(i,j).number_format=='0.0000' for i in range(2,1802) for j in range(2,23))
    assert s.freeze_panes=='B2',str(s.freeze_panes)
    checks[name]={'rows':1800,'radial_positions':21,'numeric_values':37800,
                  'max_difference_from_rounded_source':delta,'four_decimal_format':True,'freeze_panes':s.freeze_panes}
wb.close()
checks['all_passed']=True
(root/'verification/workbook_check.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False,indent=2))
