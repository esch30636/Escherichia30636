"""Read-only XLSX check; all numerical cells are compared to the same trajectory."""
from pathlib import Path
import json
import numpy as np
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
z=np.load(ROOT/'data/full_precision.npz')
w=openpyxl.load_workbook(ROOT/'result3.xlsx',data_only=False)
assert w.sheetnames==['Sheet1']
s=w.active;rows=list(s.values);a=np.array(rows[1:],dtype=float)
assert a.shape==(len(z['time_s'])-1,22)
assert np.array_equal(a[:,0],z['time_s'][1:])
assert np.array_equal(np.array(rows[0][1:],float),z['radius_cm'])
assert np.array_equal(a[:,1:],np.round(z['C'][1:],4))
assert all(s.cell(i,j).number_format=='0.0000' for i in range(2,s.max_row+1) for j in range(2,23))
assert np.isfinite(a).all() and z['maxC'][-2]>=.15 and z['maxC'][-1]<.15
assert s.freeze_panes=='B2',str(s.freeze_panes)
tab=np.loadtxt(ROOT/'data/table5.csv',delimiter=',',skiprows=1)
for row in tab:
    i=int(round(row[0]*60));assert np.max(abs(row[1:]-z['C'][i,[0,5,10,15,20]]))<1e-14
result=dict(passed=True,rows=len(a),radius_columns=21,checked_values=a[:,1:].size,
    max_rounding_difference=float(np.max(abs(a[:,1:]-z['C'][1:]))),
    first_time_s=float(a[0,0]),last_time_s=float(a[-1,0]),freeze_panes=s.freeze_panes,
    previous_max=float(z['maxC'][-2]),final_max=float(z['maxC'][-1]),table5_matches=True)
(ROOT/'verification/workbook_check.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result,indent=2))
