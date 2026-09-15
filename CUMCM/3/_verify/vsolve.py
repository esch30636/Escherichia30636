"""Independent radial-FV re-implementation of the A-problem drying model.

Deliberately differs from every candidate: uniform-in-r mesh (not the p=2 graded
mesh), arithmetic face averaging (not 3-point Gauss along the (T,C) path),
BDF with rtol=1e-10 via scipy (not the candidates' settings), no numba, no code
reuse. Used only to adjudicate, never to produce a deliverable.
"""
import numpy as np
from scipy.integrate import solve_ivp
from scipy.sparse import lil_matrix
import openpyxl

R, H, HM = 0.02, 25.0, 8e-7
ATT = r'D:/Escherichia30636/CUMCM/CUMCM2026Problems/A题/附件/附件1.xlsx'

def load_ambient():
    wb = openpyxl.load_workbook(ATT, data_only=True); ws = wb[wb.sheetnames[0]]
    rows = [r for r in ws.iter_rows(values_only=True)][1:]
    a = np.array([[float(x) for x in r[:3]] for r in rows])
    return a[:, 0], a[:, 1], a[:, 2]

TA, CT, CC = load_ambient()          # time, ambient T, ambient C

def props(t, c, ap):
    c = np.asarray(c, float)
    if ap == 2:
        return 820.*2600.*np.ones_like(c), .36*np.ones_like(c), 7e-9*np.exp(-.89/c)
    rho = 650.+128.*c; cp = 1450.+2736.*c/(1.+c)
    return rho*cp, .21+.38*c/(1.+c), 2.4e-3*np.exp(-.45/c-3850./(t+273.15))

class Model:
    def __init__(self, ap, n, tail, rtol=1e-10, max_step=60.):
        self.ap, self.n, self.tail, self.rtol, self.max_step = ap, n, np.asarray(tail,float), rtol, max_step
        self.rf = np.linspace(0., R, n+1)
        self.x = .5*(self.rf[1:]+self.rf[:-1])
        self.v = .5*(self.rf[1:]**2-self.rf[:-1]**2)
        self.dx = np.diff(self.x)
        self.half = R-self.x[-1]
    def air(self, t):
        if t >= TA[-1]: return self.tail
        return np.array([np.interp(t, TA, CT), np.interp(t, TA, CC)])
    def rhs(self, t, y):
        T, C, n = y[0::2], y[1::2], self.n
        cap, k, d = props(t, C, self.ap)
        kf = .5*(k[:-1]+k[1:]); df = .5*(d[:-1]+d[1:])
        qi = self.rf[1:n]*kf*(T[:-1]-T[1:])/self.dx
        mi = self.rf[1:n]*df*(C[:-1]-C[1:])/self.dx
        ta, ca = self.air(t); hw = self.half
        Ts = (k[-1]*T[-1]+H*hw*ta)/(k[-1]+H*hw)
        Cs = (d[-1]*C[-1]+HM*hw*ca)/(d[-1]+HM*hw)
        qo, mo = R*H*(Ts-ta), R*HM*(Cs-ca)
        o = np.empty(2*n)
        o[0::2] = (np.r_[0., qi]-np.r_[qi, qo])/(self.v*cap)
        o[1::2] = (np.r_[0., mi]-np.r_[mi, mo])/self.v
        return o
    def sparsity(self):
        a = lil_matrix((2*self.n, 2*self.n), dtype=int)
        for i in range(self.n): a[2*i:2*i+2, 2*max(0,i-1):2*min(self.n,i+2)] = 1
        return a.tocsr()
    def run(self, tmax, method='BDF', teval=None):
        y0 = np.tile([28., 2.55], self.n)
        return solve_ivp(self.rhs, (0., tmax), y0, method=method, jac_sparsity=self.sparsity(),
                         rtol=self.rtol, atol=np.tile([self.rtol*.1, self.rtol*.01], self.n),
                         max_step=self.max_step, dense_output=True)
    def fields(self, t, y):
        T, C = y[0::2], y[1::2]
        ta, ca = self.air(t); hw = self.half
        cap, k, d = props(t, C, self.ap)
        Ts = (k[-1]*T[-1]+H*hw*ta)/(k[-1]+H*hw)
        Cs = (d[-1]*C[-1]+HM*hw*ca)/(d[-1]+HM*hw)
        return T, C, Ts, Cs
    def at(self, y, r):          # r in metres
        T, C, Ts, Cs = self.fields(0., y)
        c0 = (self.x[1]**2*C[0]-self.x[0]**2*C[1])/(self.x[1]**2-self.x[0]**2)
        t0 = (self.x[1]**2*T[0]-self.x[0]**2*T[1])/(self.x[1]**2-self.x[0]**2)
        if r <= 1e-9: return t0, c0
        if r >= R-1e-12: return Ts, Cs
        return np.interp(r, self.x, T), np.interp(r, self.x, C)
    def maxC(self, t, y):
        T, C, Ts, Cs = self.fields(t, y)
        return max(C.max(), Cs, (self.x[1]**2*C[0]-self.x[0]**2*C[1])/(self.x[1]**2-self.x[0]**2))

def critical(m, tmax, method='BDF'):
    s = m.run(tmax, method=method)
    if not s.success: raise RuntimeError(s.message)
    lo, hi = 0., tmax
    f = lambda t: m.maxC(t, s.sol(t))-.15
    if f(tmax) >= 0: return None, s
    for _ in range(60):
        mid = .5*(lo+hi)
        if f(mid) < 0: hi = mid
        else: lo = mid
    return .5*(lo+hi), s
